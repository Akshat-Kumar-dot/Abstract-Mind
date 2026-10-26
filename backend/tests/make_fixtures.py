"""Builds synthetic test documents with known answers: two-column IEEE-style
and single-column LNCS-style PDFs, a Word manuscript with a title page, LaTeX
sources for several classes, and UPI receipts (image and PDF).

Run directly to (re)write tests/fixtures/, or let conftest.py do it."""
from __future__ import annotations

from pathlib import Path

import fitz

OUT = Path(__file__).parent / "fixtures"

BODY = ("Large language models are now used to read scientific documents, yet most submission systems still "
        "ask authors to type the same metadata by hand. Prior work by Smith et al. [3] studied header parsing "
        "with conditional random fields, while Lee and Wong [5] relied on font statistics. We compare layout "
        "rules with learned models on manuscripts from many templates and report where each one fails. ")


class Writer:
    def __init__(self, page: fitz.Page, font: str = "tiro", bold: str = "tibo"):
        self.page, self.font, self.bold = page, font, bold

    def text(self, x: float, y: float, s: str, size: float = 10, bold: bool = False) -> float:
        f = self.bold if bold else self.font
        self.page.insert_text((x, y), s, fontsize=size, fontname=f)
        return x + fitz.get_text_length(s, fontname=f, fontsize=size)

    def sup(self, x: float, y: float, s: str, size: float = 10) -> float:
        small = size * 0.6
        self.page.insert_text((x, y - size * 0.35), s, fontsize=small, fontname=self.font)
        return x + fitz.get_text_length(s, fontname=self.font, fontsize=small)

    def centered(self, cx: float, y: float, s: str, size: float = 10, bold: bool = False) -> None:
        f = self.bold if bold else self.font
        w = fitz.get_text_length(s, fontname=f, fontsize=size)
        self.text(cx - w / 2, y, s, size, bold)

    def para(self, x: float, y: float, width: float, s: str, size: float = 10, lead: float = 12, bold_prefix: str = "") -> float:
        words = s.split()
        line, first = "", True
        if bold_prefix:
            x0 = self.text(x, y, bold_prefix, size, bold=True)
            avail = width - (x0 - x)
        else:
            x0, avail = x, width
        for w in words:
            trial = (line + " " + w).strip()
            if fitz.get_text_length(trial, fontname=self.font, fontsize=size) > avail:
                self.text(x0 if first else x, y, line, size)
                y += lead
                line, first, avail = w, False, width
            else:
                line = trial
        if line:
            self.text(x0 if first else x, y, line, size)
            y += lead
        return y


def ieee_pdf() -> bytes:
    doc = fitz.open()
    W, H = 612, 792
    p = doc.new_page(width=W, height=H)
    w = Writer(p)
    w.centered(W / 2, 80, "Reading Submission Metadata from Unpublished", 22)
    w.centered(W / 2, 106, "Manuscripts with Layout Rules", 22)
    cols = [(W / 6, "Aarav Sharma", "dept. of Computer Science", "IIT Delhi", "New Delhi, India", "aarav.sharma@cse.iitd.ac.in"),
            (W / 2, "Meera Iyer", "School of Informatics", "University of Edinburgh", "Edinburgh, UK", "m.iyer@ed.ac.uk"),
            (5 * W / 6, "Daniel Okafor", "Microsoft Research", "Microsoft", "Redmond, USA", "dokafor@microsoft.com")]
    for cx, name, d, inst, city, mail in cols:
        w.centered(cx, 140, name, 11)
        w.centered(cx, 153, d, 10)
        w.centered(cx, 165, inst, 10)
        w.centered(cx, 177, city, 10)
        w.centered(cx, 189, mail, 10)
    left, right, cw = 54, 318, 240
    y = w.para(left, 225, cw, "This paper presents a system that extracts the title, authors, affiliations and declarations "
               "from unpublished manuscripts so that authors do not have to retype them into submission forms.", 9, 11, bold_prefix="Abstract—")
    y = w.para(left, y + 4, cw, "metadata extraction, document layout, research papers.", 9, 11, bold_prefix="Index Terms—")
    w.centered(left + cw / 2, y + 12, "I. INTRODUCTION", 10)
    y = y + 28
    for _ in range(3):
        y = w.para(left, y, cw, BODY, 10, 12) + 4
    y2 = 225
    for _ in range(3):
        y2 = w.para(right, y2, cw, BODY, 10, 12) + 4
    w.centered(right + cw / 2, y2 + 12, "II. METHOD", 10)
    y2 = w.para(right, y2 + 28, cw, BODY, 10, 12)
    w.text(W / 2 - 3, H - 30, "1", 9)

    p2 = doc.new_page(width=W, height=H)
    w2 = Writer(p2)
    y = 60
    for _ in range(4):
        y = w2.para(left, y, cw, BODY, 10, 12) + 4
    w2.centered(left + cw / 2, y + 10, "Fig. 1. Pipeline overview.", 8)
    y += 24
    w2.centered(left + cw / 2, y + 10, "TABLE I", 8)
    w2.centered(left + cw / 2, y + 20, "Results on the test set", 8)
    y += 40
    w2.centered(left + cw / 2, y, "ACKNOWLEDGMENT", 10)
    y = w2.para(left, y + 16, cw, "This work was supported by the Science and Engineering Research Board under "
                "Grant No. CRG/2021/004567. We thank Priya Raman for comments.", 10, 12)
    y2 = 60
    w2.centered(right + cw / 2, y2, "REFERENCES", 10)
    y2 += 16
    for i in range(1, 13):
        y2 = w2.para(right, y2, cw, f"[{i}] J. Smith, K. Lee, and R. Patel, \u201cOn parsing headers,\u201d in Proc. ACL, 20{10 + i}, pp. {i}-{i + 9}.", 8, 10)
    w2.text(W / 2 - 3, H - 30, "2", 9)
    return doc.tobytes()


def lncs_pdf() -> bytes:
    doc = fitz.open()
    W, H = 595, 842
    p = doc.new_page(width=W, height=H)
    w = Writer(p)
    w.centered(W / 2, 110, "A Zone-Based Method for Author Extraction", 14, bold=True)
    # Authors with superscript markers
    line = [("Alice Kumar", "1"), ("Bob Mehta", "2,*"), ("Carol Smith", "1")]
    total = 0
    parts = []
    for i, (n, m) in enumerate(line):
        sep = ", " if i < len(line) - 2 else (", and " if i == len(line) - 2 else "")
        parts.append((n, m, sep))
    x = 150
    for n, m, sep in parts:
        x = w.text(x, 140, n, 10)
        x = w.sup(x, 140, m, 10)
        if sep:
            x = w.text(x, 140, sep, 10)
    x = w.sup(140, 165, "1", 9)
    w.text(x + 1, 165, "Indian Institute of Science, Bengaluru 560012, India", 9)
    w.centered(W / 2, 177, "{alice,carol}@iisc.ac.in", 9)
    x = w.sup(160, 191, "2", 9)
    w.text(x + 1, 191, "University of Oxford, Oxford OX1 3QD, UK", 9)
    w.centered(W / 2, 203, "bob.mehta@ox.ac.uk", 9)
    y = w.para(110, 240, 375, "Submission systems ask for the same author details again and again. We split each manuscript "
               "into zones and read authors only from the header, which stops cited authors from leaking into the result.", 9, 11, bold_prefix="Abstract. ")
    y = w.para(110, y + 6, 375, "metadata · zone segmentation · author linking", 9, 11, bold_prefix="Keywords: ")
    w.text(110, y + 20, "1   Introduction", 12, bold=True)
    y = y + 38
    for _ in range(4):
        y = w.para(110, y, 375, BODY, 10, 12) + 4
    w.text(110, y + 10, "Funding. This research was funded by the Department of Science and Technology under grant number SR/FST/2020-123.", 9)
    w.text(110, y + 24, "Conflict of interest. The authors declare no competing interests.", 9)
    x = w.sup(110, 800, "*", 8)
    w.text(x + 2, 800, "Corresponding author", 8)
    p2 = doc.new_page(width=W, height=H)
    w2 = Writer(p2)
    w2.text(110, 80, "References", 12, bold=True)
    y = 100
    for i in range(1, 9):
        y = w2.para(110, y, 375, f"{i}. Garcia, M., Chen, L.: Header parsing at scale. In: Proc. ICDAR, pp. {i}0-{i}9 (20{12 + i})", 9, 11)
    return doc.tobytes()


def nature_pdf() -> bytes:
    """Nature-style layout: no Abstract/Introduction/References headings,
    authors in a block with numbered markers and ✉, notes run together in one
    page-1 footnote, numbered references in small type, bold run-in headings."""
    doc = fitz.open()
    W, H = 595, 791
    font = fitz.Font("helv")
    bold = fitz.Font("hebo")
    # Nature marks the corresponding author with ✉; fall back to * without a symbol font
    symbol, ENVELOPE = font, "*"
    for path in ("C:/Windows/Fonts/seguisym.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if Path(path).exists():
            symbol, ENVELOPE = fitz.Font(fontfile=path), "✉"
            break
    for pno in range(1, 4):
        p = doc.new_page(width=W, height=H)
        tw = fitz.TextWriter(p.rect)

        def put(x, y, s, size=8, f=font):
            tw.append((x, y), s, font=f, fontsize=size)
            return x + f.text_length(s, fontsize=size)

        def sup(x, y, s, size=8):
            tw.append((x, y - size * 0.35), s, font=font, fontsize=size * 0.6)
            return x + font.text_length(s, fontsize=size * 0.6)

        put(424, 765, f"Nature | Vol 620 | 3 August 2023 | {100 + pno}", 8)
        if pno == 1:
            put(40, 60, "Mapping glacier retreat with satellite", 24, bold)
            put(40, 88, "radar and machine learning", 24, bold)
            names = [("Asha Raman", "1,3"), ("Lukas Brandt", "1,3"), ("Wei Zhang", "2"), ("Fatima Okoye", "1")]
            x = 217
            for i, (n, m) in enumerate(names):
                x = put(x, 145, n, 8, bold)
                x = sup(x, 145, m)
                if i == 0:
                    x = put(x, 145, " " + ENVELOPE, 8, symbol)
                if i < len(names) - 2:
                    x = put(x, 145, ", ", 8, bold)
                elif i == len(names) - 2:
                    x = put(x, 145, " & ", 8, bold)
            for y, s in ((143, "https://doi.org/10.1038/s41586-023-01234-5"), (159, "Received: 2 January 2023"),
                         (175, "Accepted: 9 May 2023"), (191, "Published online: 26 July 2023"), (207, "Open access")):
                put(40, y, s, 8)
            summary = ("Glaciers are retreating faster than at any time in the observational record, yet most of the world's "
                       "glaciers are not monitored in the field. Here we show that radar images from satellites can be turned "
                       "into monthly maps of glacier fronts with an accuracy close to that of field surveys. We trained a "
                       "network on eleven thousand hand-drawn fronts and applied it to every glacier in the Himalaya, "
                       "finding that the rate of retreat has doubled since the year two thousand and that small glaciers "
                       "are the most at risk over the coming decades.")
            y = Writer(p, "helv", "hebo").para(217, 240, 340, summary, 9, 13)
            body = BODY * 2
            yl = Writer(p, "helv", "hebo").para(40, y + 30, 255, body, 8, 10.7)
            Writer(p, "helv", "hebo").para(306, y + 30, 255, body, 8, 10.7)
            tw.write_text(p)
            w2 = fitz.TextWriter(p.rect)
            x = 40
            for mark, text in (("1", "Department of Geography, University of Zurich, Zurich, Switzerland. "),
                               ("2", "Institute of Tibetan Plateau Research, Beijing, China. "),
                               ("3", "These authors contributed equally: Asha Raman, Lukas Brandt. ")):
                w2.append((x, 728 - 6 * 0.35), mark, font=font, fontsize=3.6)
                x += font.text_length(mark, fontsize=3.6)
                w2.append((x, 728), text, font=font, fontsize=6)
                x += font.text_length(text, fontsize=6)
            w2.append((40, 735), ENVELOPE, font=symbol, fontsize=6)
            w2.append((40 + symbol.text_length(ENVELOPE, fontsize=6), 735), "e-mail: a.raman@geo.uzh.ch", font=font, fontsize=6)
            w2.write_text(p)
            continue
        if pno == 2:
            y = Writer(p, "helv", "hebo").para(40, 60, 255, BODY * 2, 8, 10.7)
            put(40, y + 10, "Fig. 1 | Glacier fronts over time.", 7, bold)
            yr = 60
            for i in range(1, 16):
                put(306, yr, f"{i}. Smith, J. & Lee, K. Glacier change in High Asia. J. Glaciol. {40 + i}, 1-10 (20{10 + i}).", 6)
                yr += 8
            tw.write_text(p)
            continue
        put(40, 60, "Methods", 12, bold)
        y = Writer(p, "helv", "hebo").para(40, 80, 255, BODY, 8, 10.7)
        yr = 60
        for i in range(16, 21):
            put(306, yr, f"{i}. Garcia, M. Radar glaciology. Remote Sens. Environ. {i}, 1-9 (2019).", 6)
            yr += 8
        x = put(306, yr + 10, "Acknowledgements ", 6, bold)
        put(x, yr + 10, "We thank the Swiss Polar Institute for field support.", 6)
        x = put(306, yr + 20, "Funding ", 6, bold)
        put(x, yr + 20, "This work was funded by the Swiss National Science Foundation under grant 200021-182345.", 6)
        x = put(306, yr + 30, "Competing interests ", 6, bold)
        put(x, yr + 30, "The authors declare no competing interests.", 6)
        tw.write_text(p)
    return doc.tobytes()


def word_docx() -> bytes:
    import io

    from docx import Document
    from docx.shared import Pt

    d = Document()
    d.add_paragraph("Soil Moisture Estimation from Sentinel-1 Data over Indian Farmland", style="Title")
    p = d.add_paragraph()
    for name, mark, sep in (("Ravi Sharma", "a", ", "), ("Priya Nair", "b,*", "")):
        p.add_run(name)
        r = p.add_run(mark)
        r.font.superscript = True
        if sep:
            p.add_run(sep)
    for mark, text in (("a", "Department of Physics, University of Delhi, New Delhi 110007, India"),
                       ("b", "School of Engineering, University of Edinburgh, Edinburgh EH9 3FB, United Kingdom")):
        q = d.add_paragraph()
        r = q.add_run(mark)
        r.font.superscript = True
        q.add_run(" " + text)
    d.add_paragraph("*Corresponding author: priya.nair@ed.ac.uk")
    d.add_paragraph("Email: ravi.sharma@du.ac.in")
    d.add_page_break()
    d.add_heading("Abstract", level=1)
    d.add_paragraph("We estimate soil moisture from radar backscatter over farmland in northern India and compare "
                    "three retrieval models against field probes collected over two monsoon seasons. " * 3)
    d.add_paragraph("Keywords: soil moisture; Sentinel-1; radar; agriculture")
    d.add_heading("1. Introduction", level=1)
    for _ in range(3):
        d.add_paragraph(BODY)
    d.add_paragraph("Figure 1: Study area.")
    d.add_paragraph("Figure 2: Backscatter over time.")
    d.add_paragraph("Table 1: Field probe statistics.")
    d.add_heading("CRediT authorship contribution statement", level=1)
    d.add_paragraph("Ravi Sharma: Conceptualization, Methodology, Software. Priya Nair: Supervision, Writing – review & editing.")
    d.add_heading("Declaration of competing interest", level=1)
    d.add_paragraph("The authors declare that they have no known competing financial interests.")
    d.add_heading("Data availability", level=1)
    d.add_paragraph("Data are available at https://zenodo.org/records/1234567.")
    d.add_heading("Funding", level=1)
    d.add_paragraph("This work was funded by ISRO under grant no. ISRO/RES/3/812/2022.")
    d.add_heading("References", level=1)
    for i in range(1, 6):
        d.add_paragraph(f"Author {i}, B. (20{10 + i}). A paper about radar. Remote Sensing, {i}, 1-10.")
    for para in d.paragraphs:
        for run in para.runs:
            if run.font.size is None and para.style.name == "Normal":
                run.font.size = Pt(11)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


TEX_ACM = r"""\documentclass[sigconf]{acmart}
\begin{document}
\title{Paper Mind: Reading Manuscripts So Authors Don't Have To}
\author{Ishaan Verma}
\email{ishaan.verma@iiitd.ac.in}
\orcid{0000-0002-1825-0097}
\affiliation{\institution{IIIT Delhi}\city{New Delhi}\country{India}}
\author{Sara Lindqvist}
\authornote{Corresponding author.}
\email{sara.lindqvist@kth.se}
\affiliation{\department{School of EECS}\institution{KTH Royal Institute of Technology}\city{Stockholm}\country{Sweden}}
\begin{abstract}
We read manuscript headers directly from LaTeX source, so the author list is exact.
\end{abstract}
\keywords{metadata, LaTeX, submission forms}
\maketitle
\section{Introduction}
Text citing \cite{smith2020,lee2021} and \cite{smith2020}.
\begin{figure}\caption{One}\end{figure}
\begin{table}\caption{Two}\end{table}
\begin{acks}
This work was supported by grant number IND-2024-77.
\end{acks}
\bibliography{refs}
\end{document}
"""

TEX_IEEE = r"""\documentclass[conference]{IEEEtran}
\begin{document}
\title{Fast Header Parsing for Conference Submissions}
\author{\IEEEauthorblockN{1\textsuperscript{st} Kavya Rao}
\IEEEauthorblockA{\textit{Dept. of ECE} \\
\textit{NIT Trichy}\\
Tiruchirappalli, India \\
kavya.rao@nitt.edu}
\and
\IEEEauthorblockN{2\textsuperscript{nd} Tom Becker}
\IEEEauthorblockA{\textit{Institute of Informatics} \\
\textit{TU Munich}\\
Munich, Germany \\
becker@in.tum.de}
}
\maketitle
\begin{abstract}
Short abstract text for testing purposes only, describing nothing in particular.
\end{abstract}
\begin{IEEEkeywords}
parsing, headers
\end{IEEEkeywords}
\section{Introduction}
Body.
\begin{thebibliography}{9}
\bibitem{a} A.
\bibitem{b} B.
\bibitem{c} C.
\end{thebibliography}
\end{document}
"""

TEX_ELS = r"""\documentclass[preprint,12pt]{elsarticle}
\begin{document}
\begin{frontmatter}
\title{Groundwater Recharge in the Deccan Traps}
\author[1]{Neha Kulkarni\corref{cor1}}
\ead{neha.k@iitb.ac.in}
\author[1,2]{Rahul Deshmukh}
\cortext[cor1]{Corresponding author}
\affiliation[1]{organization={Department of Civil Engineering, IIT Bombay},city={Mumbai},postcode={400076},country={India}}
\affiliation[2]{organization={British Geological Survey},city={Nottingham},country={UK}}
\begin{abstract}
Recharge estimates for basalt aquifers are uncertain; we combine isotopes and models.
\end{abstract}
\begin{keyword}
groundwater \sep basalt \sep recharge
\end{keyword}
\end{frontmatter}
\section{Introduction}
Text.
\end{document}
"""


def upi_image(kind: str = "phonepe") -> bytes:
    import io

    from PIL import Image, ImageDraw, ImageFont

    def font(size):
        for name in ("arial.ttf", "Nirmala.ttf", "DejaVuSans.ttf"):
            for base in ("C:/Windows/Fonts/", "/usr/share/fonts/truetype/dejavu/", ""):
                try:
                    return ImageFont.truetype(base + name, size)
                except OSError:
                    continue
        return ImageFont.load_default()

    img = Image.new("RGB", (720, 1280), "white")
    d = ImageDraw.Draw(img)
    if kind == "phonepe":
        lines = [(60, 80, "Transaction Successful", 40),
                 (60, 140, "02:15 pm on 09 Oct 2025", 28),
                 (60, 240, "Paid to", 28),
                 (60, 290, "RAMESH KIRANA STORE", 34),
                 (60, 340, "ramesh.store@ybl", 28),
                 (480, 290, "₹1,250", 40),
                 (60, 460, "Transfer Details", 30),
                 (60, 520, "Transaction ID", 26),
                 (60, 560, "T2510091415223344556677", 28),
                 (60, 640, "Debited from", 26),
                 (60, 680, "XXXXXXXX4821", 28),
                 (480, 680, "₹1,250", 28),
                 (60, 730, "UTR: 528212345678", 28),
                 (60, 820, "State Bank of India", 26),
                 (60, 900, "Powered by UPI · PhonePe", 24)]
    else:
        lines = [(260, 80, "₹500", 72),
                 (200, 180, "To ANJALI GUPTA", 32),
                 (200, 230, "Completed", 28),
                 (200, 280, "9 Oct 2025, 6:42 pm", 26),
                 (60, 420, "UPI transaction ID", 26),
                 (60, 460, "412398765432", 28),
                 (60, 540, "To: ANJALI GUPTA", 26),
                 (60, 580, "anjali.g@okhdfcbank", 26),
                 (60, 660, "From: ROHIT VERMA (HDFC Bank ••••7731)", 26),
                 (60, 700, "rohit.verma@okicici", 26),
                 (60, 780, "Google transaction ID", 26),
                 (60, 820, "CICAgKDx8vWhNA", 28),
                 (60, 900, "G Pay", 30)]
    for x, y, s, size in lines:
        d.text((x, y), s, fill="black", font=font(size))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def upi_pdf() -> bytes:
    doc = fitz.open()
    p = doc.new_page(width=300, height=500)
    font = fitz.Font("helv")
    tw = fitz.TextWriter(p.rect)
    rows = [(20, 40, "Paytm", 16), (20, 70, "Paid Successfully", 14), (20, 100, "Rs. 349.00", 18),
            (20, 130, "To: Swiggy", 11), (20, 145, "UPI ID: swiggy@paytm", 10),
            (20, 170, "From: Neel Shah, ICICI Bank - 9012", 10),
            (20, 195, "UPI Ref No: 529912340001", 10), (20, 210, "Order ID: 20251009111213141516", 10),
            (20, 235, "14 Oct 2025, 08:05 PM", 10), (20, 260, "Note: dinner", 10)]
    for x, y, s, size in rows:
        tw.append((x, y), s, font=font, fontsize=size)
    tw.write_text(p)
    return doc.tobytes()


def build(out: Path = OUT) -> Path:
    out.mkdir(parents=True, exist_ok=True)
    (out / "ieee.pdf").write_bytes(ieee_pdf())
    (out / "lncs.pdf").write_bytes(lncs_pdf())
    (out / "nature.pdf").write_bytes(nature_pdf())
    (out / "manuscript.docx").write_bytes(word_docx())
    (out / "acm.tex").write_text(TEX_ACM, encoding="utf-8")
    (out / "ieee.tex").write_text(TEX_IEEE, encoding="utf-8")
    (out / "elsarticle.tex").write_text(TEX_ELS, encoding="utf-8")
    (out / "phonepe.png").write_bytes(upi_image("phonepe"))
    (out / "gpay.png").write_bytes(upi_image("gpay"))
    (out / "paytm.pdf").write_bytes(upi_pdf())
    return out


if __name__ == "__main__":
    print(build())
